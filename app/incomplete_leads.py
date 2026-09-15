"""incomplete_leads.py -- "chala to'ldirilgan" (xom/incomplete) lidlarni
topish va Excel hisobotini qurish.

2026-09, foydalanuvchi so'rovi: "agar xom to'ltirilgan bo'lsa, alohida
excel/pdf qilinsin, ularni bitta ro'yxatga olib tashlansin... mana bularni
chala to'ldiribsiz deb ma'lumotlari, ismi, nomeri bo'lsin va chala
to'ldirgani bo'lsin... lyuboy kompaniyada shunaqa bo'lsa, hamma kompaniya
uchun."

Bu -- HAR BIR kompaniya uchun umumiy (multi-tenant) mexanizm: `Lead.
full_name` yoki `Lead.phone` (ikkalasi ham -- yoki biri) bo'sh bo'lgan
lidlarni topadi (bular menejer ULARGA QO'NG'IROQ QILA OLMAYDI yoki KIM
ekanini bilmaydi -- CRM'da "ishlashning" eng minimal sharti shu ikkitasi),
va ularni bitta Excel faylga (ism/telefon ustunlari + "nima yetishmayapti"
ustuni bilan) yig'ib beradi -- menejer keyin qo'lda to'ldirishi uchun.

Format tanlovi (Excel, PDF yoki CSV) foydalanuvchi uchun farqi yo'q edi --
loyihada `openpyxl` allaqachon bog'liqlik sifatida bor (`monthly_report.py`
PDF uchun `reportlab` ishlatadi, lekin oddiy jadval uchun Excel yengilroq
va menejer to'g'ridan-to'g'ri tahrirlab, keyin CRM'ga qo'lda kiritishi
mumkin), shuning uchun Excel (.xlsx) tanlandi.
"""
from __future__ import annotations

import io
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, and_

import db


def missing_fields_for(lead: "db.Lead") -> list[str]:
    """Bitta lead uchun "yetishmayotgan" deb hisoblanadigan maydonlar
    ro'yxatini qaytaradi -- bo'sh bo'lsa lead TO'LIQ (ro'yxat bo'sh)."""
    missing = []
    if not (lead.full_name or "").strip():
        missing.append("Ism-familiya")
    if not (lead.phone or "").strip() and not (lead.phone2 or "").strip():
        missing.append("Telefon raqami")
    return missing


def find_incomplete_leads(session, company_id: "int | None" = None, limit: int = 2000) -> list["db.Lead"]:
    """Joriy (yoki berilgan) kompaniyaning ism YOKI telefon raqami bo'sh
    bo'lgan lidlarini qaytaradi, eng yangisidan boshlab. Standart
    tenant-scoping (`db.scoped_as`) chaqiruvchi tomonda allaqachon faol
    bo'lishi kerak -- bu funksiya faqat SQL filtrini quradi, kompaniyani
    o'zi tanlamaydi (shu bilan "har qanday kompaniyada" ishlaydi)."""
    q = session.query(db.Lead).filter(
        or_(
            db.Lead.full_name.is_(None), db.Lead.full_name == "",
            and_(
                or_(db.Lead.phone.is_(None), db.Lead.phone == ""),
                or_(db.Lead.phone2.is_(None), db.Lead.phone2 == ""),
            ),
        )
    )
    if company_id is not None:
        q = q.filter(db.Lead.company_id == company_id)
    return q.order_by(db.Lead.created_at.desc()).limit(limit).all()


def build_incomplete_leads_workbook(leads: list["db.Lead"], company_name: str = "") -> bytes:
    """`leads` ro'yxatidan (allaqachon `find_incomplete_leads()` orqali
    filtrlangan) Excel workbook quradi va bayt sifatida qaytaradi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Chala to'ldirilgan"

    headers = ["#", "Ism-familiya", "Telefon", "Email", "Manba (kampaniya)", "Nima yetishmayapti", "Biriktirilgan menejer", "Yaratilgan sana"]
    ws.append(headers)
    header_fill = PatternFill(start_color="FFEDE7D9", end_color="FFEDE7D9", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for i, lead in enumerate(leads, start=1):
        missing = missing_fields_for(lead)
        manager_name = ""
        if getattr(lead, "assigned_manager", None):
            manager_name = lead.assigned_manager.full_name or lead.assigned_manager.username
        ws.append([
            i,
            lead.full_name or "(bo'sh)",
            lead.phone or lead.phone2 or "(bo'sh)",
            lead.email or "",
            lead.campaign_name or lead.form_name or "",
            ", ".join(missing),
            manager_name,
            lead.created_at.strftime("%Y-%m-%d %H:%M") if lead.created_at else "",
        ])

    widths = [5, 26, 18, 26, 26, 26, 20, 16]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def incomplete_leads_filename(company_name: str = "") -> str:
    today = dt.datetime.utcnow().strftime("%Y-%m-%d")
    safe_name = "".join(c for c in (company_name or "kompaniya") if c.isalnum() or c in (" ", "-", "_")).strip() or "kompaniya"
    return f"chala-toldirilgan-lidlar-{safe_name}-{today}.xlsx".replace(" ", "_")
